"""Unified CSV/Excel loading, field mapping and validation for sales data."""

import csv
import hashlib
import json
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import TypedDict
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

REQUIRED_FIELDS = ("date", "product", "sales_amount")
FIELD_ALIASES = {
    "date": ("date", "销售日期", "日期"),
    "product": ("product", "商品名称", "商品"),
    "sales_amount": ("sales_amount", "销售额", "销售金额"),
}
ColumnMap = Mapping[str, str] | str | Path | None


class SalesRecord(TypedDict):
    date: str
    product: str
    sales_amount: str


class SalesDataError(ValueError):
    """Input data does not meet the documented sales data contract."""


@dataclass(frozen=True)
class SalesDataset:
    records: list[SalesRecord]
    source: dict


def validate_sales_fields(fieldnames: Sequence[str] | None) -> None:
    """Check canonical, nonempty and unique headers; allow extra columns."""
    if not fieldnames:
        raise SalesDataError("数据缺少表头。")
    if any(not isinstance(name, str) or not name.strip() for name in fieldnames):
        raise SalesDataError("数据表头包含空字段名或非文本字段名。")
    if len(fieldnames) != len(set(fieldnames)):
        raise SalesDataError("数据表头包含重复字段。")
    missing = set(REQUIRED_FIELDS) - set(fieldnames)
    if missing:
        raise SalesDataError(f"数据缺少必要字段：{', '.join(sorted(missing))}。")


def _read_column_map(column_map: ColumnMap) -> dict[str, str]:
    if column_map is None:
        return {}
    if isinstance(column_map, (str, Path)):
        try:
            column_map = json.loads(Path(column_map).read_text(encoding="utf-8-sig"))
        except (UnicodeError, ValueError) as exc:
            raise SalesDataError("字段映射文件必须是 UTF-8 JSON 对象。") from exc
    if (not isinstance(column_map, Mapping) or set(column_map) - set(REQUIRED_FIELDS)
            or any(not isinstance(value, str) or not value.strip() for value in column_map.values())):
        raise SalesDataError("字段映射必须是标准字段到原始表头的对象；仅支持 date、product、sales_amount。")
    return {key: value.strip() for key, value in column_map.items()}


def _resolve_headers(raw_headers: Sequence, overrides: dict[str, str]) -> tuple[list[str], dict[str, str]]:
    if not raw_headers:
        raise SalesDataError("数据缺少表头。")
    headers = [value.strip() if isinstance(value, str) else value for value in raw_headers]
    if any(not isinstance(value, str) or not value for value in headers):
        raise SalesDataError("数据表头包含空字段名或非文本字段名。")
    if len(headers) != len(set(headers)):
        raise SalesDataError("数据表头包含重复字段。")
    resolved = {}
    for field, aliases in FIELD_ALIASES.items():
        candidates = [overrides[field]] if field in overrides else [name for name in aliases if name in headers]
        if not candidates or candidates[0] not in headers:
            raise SalesDataError(f"数据缺少必要字段：{field}；请检查表头或 --column-map。")
        if len(candidates) > 1:
            raise SalesDataError(f"字段 {field} 存在多个候选列，请通过 --column-map 明确指定。")
        resolved[field] = candidates[0]
    if len(set(resolved.values())) != len(resolved):
        raise SalesDataError("多个标准字段不能映射到同一原始列。")
    return headers, resolved


def _validate_record(values: dict[str, str], location: str) -> SalesRecord:
    values = {key: value.strip() for key, value in values.items()}
    raw_date = values["date"]
    try:
        if not re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", raw_date):
            raise ValueError
        date.fromisoformat(raw_date)
    except ValueError as exc:
        raise SalesDataError(f"{location}：date 必须是有效的 YYYY-MM-DD 日期。") from exc
    if not values["product"]:
        raise SalesDataError(f"{location}：product 不能为空。")
    if not re.fullmatch(r"[0-9]+(?:\.[0-9]{1,2})?", values["sales_amount"]):
        raise SalesDataError(f"{location}：sales_amount 必须是非负金额，最多两位小数。")
    return SalesRecord(date=raw_date, product=values["product"],
                       sales_amount=format(Decimal(values["sales_amount"]), ".2f"))


def _read_csv(input_path: str | Path, overrides: dict[str, str]) -> tuple[list[SalesRecord], dict]:
    records = []
    try:
        with Path(input_path).open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle, strict=True)
            headers, mapping = _resolve_headers(next(reader, None), overrides)
            indices = {field: headers.index(name) for field, name in mapping.items()}
            for row in reader:
                if not row:
                    continue
                if len(row) != len(headers):
                    raise SalesDataError(f"第 {reader.line_num} 行：列数与表头不一致。")
                records.append(_validate_record(
                    {field: row[index] for field, index in indices.items()}, f"第 {reader.line_num} 行"
                ))
    except UnicodeError as exc:
        raise SalesDataError("CSV 必须使用 UTF-8 或 UTF-8 BOM 编码。") from exc
    except csv.Error as exc:
        raise SalesDataError("CSV 格式不正确，请检查引号和分隔符。") from exc
    if not records:
        raise SalesDataError("CSV 没有数据行。")
    return records, {"format": "csv", "sheet": None, "column_mapping": mapping}


def _excel_text(cell, field: str, location: str) -> str:
    value = cell.value
    if cell.data_type in {"f", "e"}:
        raise SalesDataError(f"{location}：{field} 含公式或 Excel 错误，请复制并粘贴为值后分析。")
    if value is None:
        return ""
    if isinstance(value, bool):
        raise SalesDataError(f"{location}：{field} 不能是布尔值。")
    if field == "date":
        if isinstance(value, datetime):
            if value.time() != time():
                raise SalesDataError(f"{location}：date 含时间，请明确转换为销售日期后分析。")
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
    if field == "sales_amount" and isinstance(value, (int, float, Decimal)):
        return format(Decimal(str(value)), "f")
    if not isinstance(value, str):
        raise SalesDataError(f"{location}：{field} 的单元格类型不合法。")
    return value


def _read_xlsx(input_path: str | Path, sheet: str | None, overrides: dict[str, str]) -> tuple[list[SalesRecord], dict]:
    from defusedxml.common import DefusedXmlException
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        workbook = load_workbook(input_path, read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, InvalidFileException, ParseError, DefusedXmlException, KeyError, ValueError, TypeError) as exc:
        raise SalesDataError("无法读取 Excel；请确认文件是有效、未加密的 .xlsx。") from exc
    try:
        names = [worksheet.title for worksheet in workbook.worksheets]
        if sheet is None:
            if len(names) != 1:
                raise SalesDataError(f"Excel 必须明确选择工作表，请用 --sheet；可用工作表：{', '.join(names)}。")
            sheet = names[0]
        if sheet not in names:
            raise SalesDataError(f"工作表不存在：{sheet}；可用工作表：{', '.join(names)}。")
        worksheet = workbook[sheet]
        # Worksheet dimension hints can be stale; inspect actual stored rows.
        worksheet.reset_dimensions()
        rows = worksheet.iter_rows()
        header_cells = list(next(rows, ()))
        while header_cells and header_cells[-1].value is None:
            header_cells.pop()
        headers, mapping = _resolve_headers([cell.value for cell in header_cells], overrides)
        indices = {field: headers.index(name) for field, name in mapping.items()}
        records = []
        for number, raw_row in enumerate(rows, start=2):
            if all(cell.value is None for cell in raw_row):
                continue
            if any(cell.value is not None for cell in raw_row[len(headers):]):
                raise SalesDataError(f"工作表 {sheet} 第 {number} 行：存在没有表头的数据列。")
            location = f"工作表 {sheet} 第 {number} 行"
            values = {field: _excel_text(raw_row[index], field, location) if index < len(raw_row) else ""
                      for field, index in indices.items()}
            records.append(_validate_record(values, location))
        if not records:
            raise SalesDataError(f"工作表 {sheet} 没有数据行。")
        return records, {"format": "xlsx", "sheet": sheet, "column_mapping": mapping}
    except SalesDataError:
        raise
    except (BadZipFile, ParseError, DefusedXmlException, KeyError, ValueError, TypeError, OverflowError) as exc:
        raise SalesDataError("Excel 内容损坏或单元格无法读取，请重新导出有效的 .xlsx。") from exc
    finally:
        workbook.close()


def load_sales_csv(input_path: str | Path, *, column_map: ColumnMap = None) -> list[SalesRecord]:
    """Preserved CSV API; return validated rows and normalized two-decimal amounts."""
    return _read_csv(input_path, _read_column_map(column_map))[0]


def load_sales_xlsx(input_path: str | Path, *, sheet: str | None = None,
                    column_map: ColumnMap = None) -> list[SalesRecord]:
    """Read one Excel worksheet into the same record schema used by CSV."""
    return _read_xlsx(input_path, sheet, _read_column_map(column_map))[0]


def load_sales_data(input_path: str | Path, *, sheet: str | None = None,
                   column_map: ColumnMap = None) -> SalesDataset:
    """Application entry: validated rows plus fingerprint, sheet and actual mapping."""
    path = Path(input_path).resolve()
    fingerprint = hashlib.sha256(path.read_bytes()).hexdigest()
    overrides = _read_column_map(column_map)
    if path.suffix.lower() == ".csv":
        if sheet is not None:
            raise SalesDataError("CSV 不包含工作表，请移除 --sheet。")
        records, metadata = _read_csv(path, overrides)
    elif path.suffix.lower() == ".xlsx":
        records, metadata = _read_xlsx(path, sheet, overrides)
    else:
        raise SalesDataError("仅支持 .csv 和 .xlsx 文件；旧版 .xls 请先另存为 .xlsx。")
    if fingerprint != hashlib.sha256(path.read_bytes()).hexdigest():
        raise SalesDataError("读取期间输入文件发生变化，请重新运行。")
    return SalesDataset(records, {"name": path.name, "path": str(path), "sha256": fingerprint, **metadata})
